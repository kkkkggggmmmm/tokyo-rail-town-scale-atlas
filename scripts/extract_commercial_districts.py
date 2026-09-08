"""Extract the official 2021 Economic Census location table 2; no geometry inference.

Run: python extract_commercial_districts.py INPUT.xlsx OUTPUT.json
Source columns and suppression rules are verified against the workbook / ricchi_riyou.pdf.
"""
import hashlib
import json
import sys
from collections import Counter
from pathlib import Path

import openpyxl

SOURCE = {
    "id": "estat2021_ricchi_table2",
    "title": "総務省・経済産業省 令和3年経済センサス-活動調査 立地環境特性編 第2表",
    "table_id": "0004015880",
    "file_stat_infid": "000040186981",
    "url": "https://www.e-stat.go.jp/stat-search/files?layout=dataset&stat_infid=000040186981",
    "download_url": "https://www.e-stat.go.jp/stat-search/file-download?statInfId=000040186981&fileKind=0",
    "definition_url": "https://www.stat.go.jp/data/e-census/2021/kekka/pdf/ricchi_riyou.pdf",
    "terms_url": "https://www.e-stat.go.jp/terms-of-use",
    "license": "CC BY 4.0 compatible e-Stat terms; numerical data freely reusable; source and processing attribution required",
    "published_date": "2024-06-25",
    "retrieved_date": "2026-09-08",
    "sales_reference_period": "2020-01-01/2020-12-31",
    "establishments_reference_date": "2021-06-01",
    "sales_unit": "百万円",
    "scope": "公式商業集積地区（概ね商店街、ショッピングセンター、駅ビル等）。駅圏・市区町村全体・独自centerではない。",
    "sales_definition": "小売・飲食サービス・生活関連サービスの売上（収入）金額。GDP（付加価値）や全産業売上ではなく、小売の年間商品販売額とも集計項目が異なる。",
    "coverage_notes": [
        "1都3県の地区明細のみ。都道府県・市区町村小計を重複加算しない。",
        "飲食サービスは中分類76+77。生活関連サービスは78+79からリネンサプライと火葬・墓地管理を除く。娯楽業80や宿泊業75は含まない。",
        "売場面積は法人小売業のみ。事務所床面積ではない。",
        "2020年のコロナ期売上であり、現在の規模を表す値に換算しない。",
        "地区の地理的範囲はこのExcelにない。名称一致だけで駅／centerの境界一致と認定しない。",
        "2026-09-08に統計局正誤情報を確認。立地環境特性編第2表を対象とする訂正告知は見当たらず。",
    ],
    "corrections_url": "https://www.stat.go.jp/data/seigo/e-census/2021/index.html",
}


def observation(raw, coordinate, unit):
    if isinstance(raw, (int, float)):
        return {"value": raw, "status": "below_rounding_unit" if raw == 0 and unit == "百万円" else "observed", "raw": raw, "cell": coordinate}
    status = {"x": "suppressed", "X": "suppressed", "-": "not_applicable", "": "not_reported", None: "not_reported"}.get(raw)
    if status is None:
        raise ValueError(f"Unknown source token {raw!r} in {coordinate}")
    return {"value": None, "status": status, "raw": raw, "cell": coordinate}


def extract(path):
    book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheet = book["k002"]
    source = dict(SOURCE, sha256=hashlib.sha256(Path(path).read_bytes()).hexdigest())
    headers = list(sheet.iter_rows(min_row=10, max_row=12, values_only=True))
    assert headers[0][11] == "I2_小売業計" and headers[1][11] == "売上（収入）金額" and headers[2][11] == "百万円"
    assert headers[0][16] == "M2_飲食サービス業計" and headers[2][16] == "百万円"
    assert headers[0][20] == "N_生活関連サービス業計" and headers[2][20] == "百万円"
    municipalities = {}
    rows = []
    specs = {
        "retail_establishments": (9, "事業所"), "retail_employees": (11, "人"),
        "retail_sales_million_yen": (12, "百万円"), "retail_sales_floor_sqm": (13, "㎡"),
        "food_service_establishments": (14, "事業所"), "food_service_employees": (16, "人"),
        "food_service_sales_million_yen": (17, "百万円"),
        "personal_service_establishments": (18, "事業所"), "personal_service_employees": (20, "人"),
        "personal_service_sales_million_yen": (21, "百万円"),
    }
    for row_number, row in enumerate(sheet.iter_rows(values_only=True), 1):
        if not row[1] or str(row[1])[:2] not in {"11", "12", "13", "14"} or not row[0]:
            continue
        code = str(row[0]).split("_", 1)[0]
        if not row[3]:
            if row[2] and str(row[4]).endswith("計"):
                municipalities[code[:5]] = row[4][:-1]
            continue
        metrics = {key: observation(row[column-1], f"{openpyxl.utils.get_column_letter(column)}{row_number}", unit) for key, (column, unit) in specs.items()}
        values = [metrics[k]["value"] for k in ("retail_sales_million_yen", "food_service_sales_million_yen", "personal_service_sales_million_yen")]
        # No reconstruction of suppressed totals and no zero filling of '-'.
        total = sum(values) if all(v is not None for v in values) else None
        rows.append({"source_district_id": code, "source_id": source["id"], "name": row[4],
                     "prefecture_code": code[:2], "municipality_code": code[:5],
                     "municipality_name": municipalities.get(code[:5]),
                     "location_type": row[5], "sheet": "k002", "row": row_number,
                     "metrics": metrics,
                     "three_sector_sales_million_yen": total,
                     "three_sector_sales_status": "derived_sum_of_published_components" if total is not None else "not_calculated_missing_component"})
    assert len(rows) == 2706 and len({r["source_district_id"] for r in rows}) == len(rows)
    return {"source": source, "counts": dict(Counter(r["prefecture_code"] for r in rows)), "districts": rows}


if __name__ == "__main__":
    output = extract(sys.argv[1])
    Path(sys.argv[2]).write_text(json.dumps(output, ensure_ascii=False, separators=(",", ":")) + "\n")
    print(json.dumps({"districts": len(output["districts"]), "counts": output["counts"], "output": sys.argv[2]}, ensure_ascii=False))
